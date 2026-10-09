import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import MagicMock

import openpyxl
import pytest
from cryptography.fernet import Fernet
from openpyxl.comments import Comment
from openpyxl.styles import Font
from app import create_app
from logic import InputError, TransferGrades

@pytest.fixture
def app(tmp_path):
    return create_app({'TESTING': True, 'SECRET_KEY': 's'*40,
        'TOKEN_ENCRYPTION_KEY': Fernet.generate_key(), 'PUBLIC_BASE_URL': 'https://school.test',
        'TEACHER_DOMAIN': 'school.test',
        'STUDENT_DOMAINS': 'students.test', 'DATABASE': str(tmp_path/'state.sqlite3')})

def login(app, email='teacher@school.test'):
    client = app.test_client()
    sid = app.extensions['store'].create_session({'email': email, 'hd': 'school.test', 'sub': email, 'credentials': {'token':'TOP_SECRET'}}, 3600)
    with client.session_transaction() as session:
        session.update(sid=sid, csrf='test-csrf')
    return client, sid

def post(client, route, **kwargs):
    return client.post(route, headers={'X-CSRF-Token': 'test-csrf', 'Origin':'https://school.test'}, **kwargs)

def workbook(rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Notas'
    for row in rows or [['Nombre','Correo','Nota'],['Ana','ana@students.test',7],['Luis','luis@students.test',8]]:
        ws.append(row)
    ws['C2'].comment = Comment('Comentario individual', 'Autor privado')
    ws['C2'].font = Font(bold=True, color='FF0000')
    ws['A2'].hyperlink = 'https://private.invalid/'
    wb.create_sheet('Secreto').append(['Todos los alumnos'])
    wb['Secreto'].sheet_state = 'hidden'
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    return data

CONFIG={'numero_cabecera':'1','letra_columna_nombre':'A','letra_columna_correo':'B','nombre_hoja':'Notas'}

def test_individual_workbooks():
    items=TransferGrades(None).copy_grades(workbook(),CONFIG,{'students.test'})
    assert len(items)==2
    wb=openpyxl.load_workbook(items[0]['content'])
    assert wb.sheetnames==['Evaluación']
    assert wb.active.max_row==2
    assert wb.active['A2'].value=='Ana'
    assert wb.active['A2'].hyperlink is None
    assert wb.active['C2'].comment.author=='Docente'
    assert wb.active['C2'].font.bold
    assert 'Luis' not in str(list(wb.active.values))

@pytest.mark.parametrize('rows',[
    [['N','C','V'],['Ana','ana@external.test',1]],
    [['N','C','V'],['Ana','ana@students.test',1],['Otra','ana@students.test',2]],
    [['N','C','V'],['Ana','',1]],
    [['N','C','V'],['Ana','ana@students.test','=1+1']],
])
def test_bad_source_rejected(rows):
    with pytest.raises(InputError):
        TransferGrades(None).copy_grades(workbook(rows),CONFIG,{'students.test'})

@pytest.mark.parametrize('change',[{'numero_cabecera':'51'},{'letra_columna_nombre':'../x'},{'letra_columna_nombre':'B'},
    {'plantilla_cabecera':'template'},{'lista_otras_hoja':'Secreto'},{'nombre_hoja':'Missing'}])
def test_bad_config_rejected(change):
    with pytest.raises(InputError):
        TransferGrades(None).copy_grades(workbook(),CONFIG|change,{'students.test'})

def test_literal_not_formula():
    source=workbook()
    wb=openpyxl.load_workbook(source)
    wb['Notas']['C2']='=malicious()'
    wb['Notas']['C2'].data_type='s'
    source=io.BytesIO(); wb.save(source); source.seek(0)
    item=TransferGrades(None).copy_grades(source,CONFIG,{'students.test'})[0]
    out=openpyxl.load_workbook(item['content'])
    assert out.active['C2'].data_type=='s'

def test_concurrent_processing_isolated():
    def run(name):
        return TransferGrades(None).copy_grades(workbook([['N','C','V'],[name,name+'@students.test',1]]),CONFIG,{'students.test'})
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(run,['uno','dos']))
    assert results[0][0]['name']=='uno' and results[1][0]['name']=='dos'

def test_cookie_and_database_no_plaintext(app):
    client,sid=login(app)
    cookie=client.get_cookie('trasladar_session').value
    assert 'TOP_SECRET' not in cookie
    with client.session_transaction() as session:
        assert set(session)=={'sid','csrf'}
    assert b'TOP_SECRET' not in open(app.config['DATABASE'],'rb').read()
    assert b'teacher@school.test' not in open(app.config['DATABASE'],'rb').read()
    assert sid.encode() not in open(app.config['DATABASE'],'rb').read()

def test_logout_revokes_copied_cookie(app):
    client,sid=login(app)
    cookie=client.get_cookie('trasladar_session').value
    assert post(client,'/logout').status_code==302
    other=app.test_client(); other.set_cookie('trasladar_session',cookie)
    assert post(other,'/generate').status_code==401
    assert app.extensions['store'].session(sid) is None

def test_csrf_origin_and_anonymous(app):
    client,_=login(app)
    assert client.post('/generate').status_code==403
    assert client.post('/generate',headers={'X-CSRF-Token':'test-csrf','Origin':'https://evil.test'}).status_code==403
    client,_=login(app, 'external@other.test')
    assert post(client,'/generate').status_code==403

def test_expiry_and_purge(app):
    store=app.extensions['store']; sid=store.create_session({'email':'expired'},-1)
    job=store.create_job(sid)
    assert store.session(sid) is None
    store.purge()
    assert store.job(job,sid) is None

def test_owner_bound_job_and_one_claim(app):
    a,sid=login(app); b,_=login(app,'other@school.test')
    store=app.extensions['store']; job=store.create_job(sid)
    store.save_job(job,{'items':[]},'ready')
    assert post(b,'/notify',json={'batch_id':job,'confirmed':True}).status_code==403
    with ThreadPoolExecutor(max_workers=2) as pool:
        result=list(pool.map(lambda _:store.claim(job,sid),range(2)))
    assert sorted(result)==[False,True]
    assert post(a,'/notify',json={'batch_id':job,'confirmed':True}).status_code==409

@pytest.mark.parametrize('body',[{'items':[{'email':'evil','folder_id':'victim'}]}, [], {'batch_id':'x','confirmed':False}])
def test_browser_cannot_supply_recipients(app,body):
    client,_=login(app)
    assert post(client,'/notify',json=body).status_code==400

def test_headers_privacy_escaping(app):
    app.config['PRIVACY_CONTROLLER']='<script>alert(1)</script>'
    response=app.test_client().get('/privacidad')
    assert b'&lt;script&gt;' in response.data
    assert response.headers['Cache-Control']=='no-store'
    assert "frame-ancestors 'none'" in response.headers['Content-Security-Policy']

def test_pending_oauth_cannot_generate(app):
    client=app.test_client()
    sid=app.extensions['store'].create_session({'oauth': {'state':'s'}},600)
    with client.session_transaction() as session: session.update(sid=sid,csrf='test-csrf')
    assert post(client,'/generate').status_code==401
    assert b'Conectar con Google Drive' in client.get('/').data

def test_permission_safety_and_reader():
    drive=MagicMock()
    drive.permissions().list().execute.return_value={'permissions':[{'type':'user','role':'owner'}]}
    transfer=TransferGrades(drive)
    transfer.share({'file_id':'own','email':'ana@students.test'})
    args=drive.permissions().create.call_args.kwargs
    assert args['body']['role']=='reader' and args['fileId']=='own'
    drive.permissions().list().execute.return_value={'permissions':[{'type':'anyone','role':'reader'}]}
    drive.permissions().create.reset_mock()
    with pytest.raises(InputError): transfer.share({'file_id':'own','email':'ana@students.test'})
    drive.permissions().create.assert_not_called()

def test_ambiguous_source_and_query_escaping():
    drive=MagicMock(); drive.files().list().execute.return_value={'files':[{'id':'1'},{'id':'2'}]}
    with pytest.raises(InputError): TransferGrades(drive).download("a' or x",True)
    assert "a\\' or x" in drive.files().list.call_args.kwargs['q']

def oauth_setup(app, monkeypatch, overrides=None):
    client=app.test_client(); store=app.extensions['store']
    sid=store.create_session({'oauth':{'state':'state','nonce':'nonce','verifier':'verifier'}},600)
    with client.session_transaction() as session: session['sid']=sid
    oauth=MagicMock(); oauth.client_config={'client_id':'client'}
    oauth.credentials.to_json.return_value=json.dumps({'token':'secret','client_id':'client','client_secret':'secret','refresh_token':'refresh'})
    monkeypatch.setattr('app.Flow.from_client_secrets_file',lambda *a,**k:oauth)
    claims={'email':'teacher@school.test','hd':'school.test','sub':'subject','nonce':'nonce','email_verified':True}
    claims.update(overrides or {})
    monkeypatch.setattr('app.id_token.verify_oauth2_token',lambda *a,**k:claims)
    return client,oauth,sid

@pytest.mark.parametrize('claims',[{'hd':'evil.test'},{'email_verified':False},{'nonce':'evil'}, {'sub':''}, {'email':'pupil@students.test'}, {'email':'user@sub.school.test'}, {'email':'user@school.test.evil.test'}, {'hd':None}, {'email':'user@gmail.com'}])
def test_oauth_rejects_bad_identity(app,monkeypatch,claims):
    client,oauth,sid=oauth_setup(app,monkeypatch,claims)
    assert client.get('/oauth2callback?state=state&code=code').status_code==403
    assert app.extensions['store'].session(sid) is None

def test_oauth_state_and_rotation(app,monkeypatch):
    client,oauth,sid=oauth_setup(app,monkeypatch)
    assert client.get('/oauth2callback?state=evil').status_code==403
    oauth.fetch_token.assert_not_called()
    assert client.get('/oauth2callback?state=state&code=code').status_code==302
    with client.session_transaction() as session:
        new=app.extensions['store'].session(session['sid'])
        assert session['sid'] != sid
        assert new['credentials']['refresh_token']==''
    assert app.extensions['store'].session(sid) is None
    assert client.get('/oauth2callback?state=state&code=code').status_code==403

def pipeline(app, monkeypatch):
    client,sid=login(app)
    from google.oauth2.credentials import Credentials
    monkeypatch.setattr('app.Credentials.from_authorized_user_info',lambda info:Credentials('fake-token'))
    monkeypatch.setattr('app.build',lambda *a,**k:MagicMock())
    monkeypatch.setattr(TransferGrades,'download',lambda *a:workbook())
    monkeypatch.setattr(TransferGrades,'new_folder',lambda self,job:'new-private-folder')
    monkeypatch.setattr(TransferGrades,'upload',lambda self,item,*a:'id-'+item['name'])
    monkeypatch.setattr(TransferGrades,'assert_private',lambda *a:None)
    sent=[]
    monkeypatch.setattr(TransferGrades,'share',lambda self,item:sent.append(item.copy()))
    return client,sid,sent

def test_generate_review_notify_server_manifest(app,monkeypatch):
    client,sid,sent=pipeline(app,monkeypatch)
    response=post(client,'/generate',data=CONFIG|{'nombre_excel_notas':'source','reviewed_source':'yes'})
    assert response.status_code==200
    data=response.get_json(); job=data['batch_id']
    assert len(data['items'])==2
    assert app.extensions['store'].job(job,sid)['state']=='ready'
    data['items'][0]['email']='attacker@students.test'
    assert post(client,'/notify',json={'batch_id':job,'confirmed':True,'items':data['items']}).status_code==400
    assert not sent
    assert post(client,'/notify',json={'batch_id':job,'confirmed':True}).status_code==200
    assert [item['email'] for item in sent]==['ana@students.test','luis@students.test']
    assert app.extensions['store'].job(job,sid)['state']=='done'
    assert post(client,'/notify',json={'batch_id':job,'confirmed':True}).status_code==409
    assert len(sent)==2

def test_partial_notify_requires_manual_review(app,monkeypatch):
    client,sid,sent=pipeline(app,monkeypatch)
    job=post(client,'/generate',data=CONFIG|{'nombre_excel_notas':'source','reviewed_source':'yes'}).get_json()['batch_id']
    def share(self,item):
        if sent: raise RuntimeError('sensitive-google-response')
        sent.append(item.copy())
    monkeypatch.setattr(TransferGrades,'share',share)
    response=post(client,'/notify',json={'batch_id':job,'confirmed':True})
    assert response.status_code==409 and b'sensitive-google-response' not in response.data
    stored=app.extensions['store'].job(job,sid)
    assert stored['state']=='review_required'
    assert [item['status'] for item in stored['items']]==['shared','pending']
    assert post(client,'/notify',json={'batch_id':job,'confirmed':True}).status_code==409

def test_generation_failure_not_shareable(app,monkeypatch):
    client,sid,_=pipeline(app,monkeypatch)
    count=[]
    def upload(self,item,*a):
        if count: raise RuntimeError('private-exception')
        count.append(True)
        return 'first-file'
    monkeypatch.setattr(TransferGrades,'upload',upload)
    response=post(client,'/generate',data=CONFIG|{'nombre_excel_notas':'source','reviewed_source':'yes'})
    assert response.status_code==400 and b'private-exception' not in response.data
    job=response.get_json()['batch_id']
    assert app.extensions['store'].job(job,sid)['state']=='failed'
    assert post(client,'/notify',json={'batch_id':job,'confirmed':True}).status_code==409

def test_reject_generation_without_review_and_old_target(app,monkeypatch):
    client,_,_=pipeline(app,monkeypatch)
    assert post(client,'/generate',data=CONFIG).status_code==400
    assert post(client,'/generate',data=CONFIG|{'reviewed_source':'yes','target_folder_id':'old'}).status_code==400

def test_new_folder_never_reuses_old_or_shares_parent():
    drive=MagicMock()
    drive.files().create().execute.return_value={'id':'new'}
    drive.permissions().list().execute.return_value={'permissions':[{'type':'user','role':'owner'}]}
    assert TransferGrades(drive).new_folder('batch')=='new'
    assert drive.files().create.call_args.kwargs['body']['parents']==['root']
    drive.files().list.assert_not_called()
    drive.permissions().create.assert_not_called()

def test_revoke_cli(app):
    client,sid=login(app)
    assert app.test_cli_runner().invoke(args=['revoke-sessions']).exit_code==0
    assert app.extensions['store'].session(sid) is None


@pytest.mark.parametrize('email', ['new.employee@school.test', 'NEW.EMPLOYEE@SCHOOL.TEST'])
def test_corporate_employee_without_individual_list(app, monkeypatch, email):
    assert 'TEACHER_EMAILS' not in app.config
    client, _, _ = oauth_setup(app, monkeypatch, {'email': email})
    assert client.get('/oauth2callback?state=state&code=code').status_code == 302
    with client.session_transaction() as session:
        identity = app.extensions['store'].session(session['sid'])
    assert identity['email'] == email.lower()
    assert identity['hd'] == 'school.test'
    assert b'Conectar con Google Drive' not in client.get('/').data


def test_domain_change_revokes_existing_session(app):
    client, sid = login(app)
    app.config['TEACHER_DOMAIN'] = 'new-school.test'
    client.get('/')
    assert app.extensions['store'].session(sid) is None


def test_legacy_session_requires_new_google_login(app):
    client = app.test_client()
    sid = app.extensions['store'].create_session({'email': 'teacher@school.test', 'sub': 'subject'}, 3600)
    with client.session_transaction() as session:
        session.update(sid=sid, csrf='test-csrf')
    assert b'Conectar con Google Drive' in client.get('/').data
    assert app.extensions['store'].session(sid) is None
