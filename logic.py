"""Individual workbooks in memory; reuse checked student folders and update existing files."""
import io
import re
import zipfile
from copy import copy
import openpyxl
from openpyxl.comments import Comment
from openpyxl.utils import column_index_from_string
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
SHEET = 'application/vnd.google-apps.spreadsheet'
FOLDER = 'application/vnd.google-apps.folder'

class InputError(ValueError):
    pass

def quote(value):
    return value.replace('\\', '\\\\').replace("'", "\\'")

class TransferGrades:
    def __init__(self, drive):
        self.drive = drive

    def download(self, name, is_sheet):
        if not name or len(name) > 200:
            raise InputError('Indica un nombre de archivo válido.')
        if not is_sheet and not name.lower().endswith('.xlsx'):
            name += '.xlsx'
        kind = SHEET if is_sheet else XLSX
        result = self.drive.files().list(
            q=f"name = '{quote(name)}' and mimeType = '{kind}' and trashed = false",
            fields='files(id),nextPageToken', pageSize=2).execute()
        if len(result.get('files', [])) != 1 or result.get('nextPageToken'):
            raise InputError('Debe existir un único archivo con ese nombre y tipo. Renómbralo si hay duplicados.')
        file_id = result['files'][0]['id']
        request = (self.drive.files().export_media(fileId=file_id, mimeType=XLSX)
                   if is_sheet else self.drive.files().get_media(fileId=file_id))
        content = io.BytesIO()
        downloader = MediaIoBaseDownload(content, request, chunksize=1024 * 1024)
        done = False
        while not done:
            _, done = downloader.next_chunk()
            if content.tell() > 20 * 1024 * 1024:
                raise InputError('El archivo supera el límite de 20 MB.')
        content.seek(0)
        return content

    def copy_grades(self, source, config, domains, template=None):
        if config.get('lista_otras_hoja'):
            raise InputError('No se admiten hojas adicionales porque podrían incluir datos de otros alumnos.')
        if config.get('plantilla_cabecera') and template is None:
            raise InputError('No se ha podido cargar la plantilla de cabecera.')
        try:
            header = int(config.get('numero_cabecera', 0))
            name_col = column_index_from_string(config.get('letra_columna_nombre', '').upper())
            email_col = column_index_from_string(config.get('letra_columna_correo', '').upper())
            if not 0 <= header <= 50 or name_col == email_col:
                raise ValueError()
        except (ValueError, TypeError):
            raise InputError('Revisa las columnas y las filas de cabecera (0–50).') from None
        with zipfile.ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise InputError('El contenido descomprimido supera 100 MB.')
        source.seek(0)
        original = openpyxl.load_workbook(source, data_only=False, keep_links=False)
        source.seek(0)
        values = openpyxl.load_workbook(source, data_only=True, keep_links=False)
        try:
            sheet = config.get('nombre_hoja', '').strip() or original.sheetnames[0]
            if sheet not in original.sheetnames:
                raise InputError('No existe la pestaña indicada.')
            ws, cached = original[sheet], values[sheet]
            if ws.max_row > 2050 or ws.max_column > 200 or max(name_col, email_col) > ws.max_column:
                raise InputError('Revisa las columnas: máximo 2050 filas y 200 columnas.')
            template_bytes = None
            if template is not None:
                with zipfile.ZipFile(template) as archive:
                    if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                        raise InputError('La plantilla descomprimida supera 100 MB.')
                template.seek(0)
                template_bytes = template.read()
                check = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=False, keep_links=False)
                try:
                    template_ws = check[sheet] if sheet in check.sheetnames else check.active
                    if template_ws.max_column > 200:
                        raise InputError('La plantilla supera el máximo de 200 columnas.')
                    for row_cells in template_ws.iter_rows(min_row=header + 1):
                        if any(cell.value is not None or cell.comment or cell.hyperlink for cell in row_cells):
                            raise InputError('La plantilla solo puede contener datos en las filas de cabecera.')
                finally:
                    check.close()
            results, seen, total_size = [], set(), 0
            for row in range(header + 1, ws.max_row + 1):
                name = str(cached.cell(row, name_col).value or '').strip()
                email = str(cached.cell(row, email_col).value or '').strip().lower()
                if not name and not email:
                    continue
                if not name or not re.fullmatch(r'[^\s@,;<>]+@[^\s@,;<>]+', email) or email.rsplit('@', 1)[-1] not in domains:
                    raise InputError(f'Revisa nombre y correo institucional en la fila {row}.')
                if email in seen:
                    raise InputError('Hay destinatarios duplicados. Revisa el archivo antes de repartir.')
                seen.add(email)
                if template_bytes is not None:
                    wb = openpyxl.load_workbook(io.BytesIO(template_bytes), data_only=False, keep_links=False)
                    out = wb[sheet] if sheet in wb.sheetnames else wb.active
                    for other in list(wb.worksheets):
                        if other is not out:
                            wb.remove(other)
                    out.title = 'Evaluación'
                    # The template is only a layout/header source. Remove links
                    # and normalise comment authors before adding student data.
                    for row_cells in out.iter_rows():
                        for cell in row_cells:
                            if isinstance(cell, openpyxl.cell.cell.MergedCell):
                                continue
                            cell.hyperlink = None
                            if cell.comment:
                                cell.comment = Comment(cell.comment.text, 'Docente')
                    source_rows = [(header + 1, row)]
                else:
                    wb = openpyxl.Workbook()
                    out = wb.active
                    out.title = 'Evaluación'
                    source_rows = list(enumerate([*range(1, header + 1), row], 1))
                # Fresh workbook or reviewed header template: one student row.
                for target, origin in source_rows:
                    for col in range(1, ws.max_column + 1):
                        old = ws.cell(origin, col)
                        val = cached.cell(origin, col).value
                        if old.data_type == 'f' and val is None:
                            raise InputError('Hay fórmulas sin resultado guardado. Recalcula y guarda el original.')
                        cell = out.cell(target, col, val)
                        if isinstance(val, str):
                            cell.data_type = 's'
                        if old.has_style:
                            for attr in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
                                setattr(cell, attr, copy(getattr(old, attr)))
                        if old.comment:
                            cell.comment = Comment(old.comment.text, 'Docente')
                data = io.BytesIO()
                wb.save(data)
                wb.close()
                total_size += data.tell()
                if total_size > 50 * 1024 * 1024 or len(results) >= 500:
                    raise InputError('Divide el lote: máximo 500 destinatarios y 50 MB de copias.')
                data.seek(0)
                source_name = config.get('nombre_excel_notas', '').strip()
                if source_name.lower().endswith('.xlsx'):
                    source_name = source_name[:-5]
                selected_sheet = config.get('nombre_hoja', '').strip()
                folder_name = f'{name}__{source_name}' if selected_sheet else name
                file_name = selected_sheet if selected_sheet else f'{name}__{source_name}'
                results.append({'name': name, 'email': email, 'content': data,
                    'folder_name': folder_name, 'file_name': file_name})
            if not results:
                raise InputError('No hay filas de alumnado para repartir.')
            return results
        finally:
            original.close()
            values.close()

    def permissions(self, file_id):
        result, token = [], None
        while True:
            page = self.drive.permissions().list(fileId=file_id, pageToken=token,
                fields='permissions(id,type,role,emailAddress),nextPageToken').execute()
            result.extend(page.get('permissions', []))
            token = page.get('nextPageToken')
            if not token:
                return result

    def assert_private(self, file_id, email=None, roles=('reader',)):
        for p in self.permissions(file_id):
            if p.get('role') == 'owner' and p.get('type') == 'user':
                continue
            if email and p.get('type') == 'user' and p.get('emailAddress', '').lower() == email and p.get('role') in roles:
                continue
            raise InputError(f'El elemento de Drive {file_id} tiene un permiso no admitido ({p.get("type", "desconocido")}: {p.get("role", "desconocido")}). Revisa Compartir: solo propietario y alumno destinatario con el permiso previsto.')

    def upload(self, item, folder_id, convert):
        name = re.sub(r'[\\/\x00-\x1f]', '_', item['file_name']).strip()
        if not name:
            raise InputError('El nombre del documento no es válido.')
        name = name if convert else name + '.xlsx'
        kind = SHEET if convert else XLSX
        self.assert_private(folder_id, item['email'])
        existing = self.find_child(name, folder_id, kind)
        media = MediaIoBaseUpload(item['content'], mimetype=XLSX)
        if existing:
            self.normalize_existing_access(existing, item['email'], 'commenter' if convert else 'reader')
            file = self.drive.files().update(fileId=existing,
                body={'mimeType': kind}, media_body=media, fields='id').execute()
        else:
            file = self.drive.files().create(body={
                'name': name, 'mimeType': kind, 'parents': [folder_id]},
                media_body=media, fields='id').execute()
        return file['id']

    def find_child(self, name, parent_id, kind):
        result = self.drive.files().list(
            q=f"name = '{quote(name)}' and '{quote(parent_id)}' in parents and trashed = false",
            fields='files(id,mimeType),nextPageToken', pageSize=2).execute()
        matches = result.get('files', [])
        if len(matches) > 1 or result.get('nextPageToken'):
            raise InputError(f'Hay varios elementos llamados {name!r} en el destino. Revisa los duplicados en Drive.')
        if matches and matches[0].get('mimeType') != kind:
            raise InputError(f'El elemento {name!r} tiene un tipo distinto al esperado. Revisa el destino.')
        return matches[0]['id'] if matches else None

    def resolve_folder_path(self, path):
        path = (path or '').strip()
        if not path or path.lower() == 'root':
            return 'root'
        if len(path) > 500:
            raise InputError('La ruta de destino es demasiado larga.')
        parts = [part.strip() for part in path.replace('\\', '/').split('/') if part.strip()]
        if not parts or len(parts) > 10:
            raise InputError('La ruta de destino debe tener entre 1 y 10 carpetas.')
        parent_id = 'root'
        for part in parts:
            if part in ('.', '..') or len(part) > 100 or any(ord(char) < 32 for char in part):
                raise InputError('La ruta de destino contiene un nombre de carpeta no válido.')
            result = self.drive.files().list(
                q=(f"name = '{quote(part)}' and mimeType = '{FOLDER}' and "
                   f"'{parent_id}' in parents and trashed = false"),
                fields='files(id),nextPageToken', pageSize=2).execute()
            matches = result.get('files', [])
            if len(matches) > 1 or result.get('nextPageToken'):
                raise InputError(f'Hay varias carpetas llamadas {part!r} en la ruta. Renombra una de ellas.')
            if matches:
                parent_id = matches[0]['id']
            else:
                folder = self.drive.files().create(body={
                    'name': part, 'mimeType': FOLDER, 'parents': [parent_id]}, fields='id').execute()
                parent_id = folder['id']
            # A destination shared with other people would make every child
            # inherit access, so it cannot be used for individual grades.
            self.assert_private(parent_id)
        return parent_id

    def new_student_folder(self, name, parent_id, email=None):
        safe_name = re.sub(r'[\\/\x00-\x1f]', '_', name).strip()[:350]
        if not safe_name:
            raise InputError('No se puede crear la carpeta del alumno: el nombre no es válido.')
        existing = self.find_child(safe_name, parent_id, FOLDER)
        if existing:
            self.normalize_existing_access(existing, email, 'reader')
            return existing
        folder = self.drive.files().create(body={
            'name': safe_name, 'mimeType': FOLDER, 'parents': [parent_id]}, fields='id').execute()
        self.assert_private(folder['id'])
        return folder['id']

    def normalize_existing_access(self, file_id, email, role):
        # Only downgrade the intended student's old editing access. Never alter third-party access.
        self.assert_private(file_id, email, ('reader', 'commenter', 'writer'))
        for permission in self.permissions(file_id):
            if permission.get('type') == 'user' and permission.get('emailAddress', '').lower() == email and permission.get('role') != 'owner':
                if permission.get('role') == 'writer' or (role == 'reader' and permission.get('role') == 'commenter'):
                    self.drive.permissions().update(fileId=file_id, permissionId=permission['id'], body={'role':role}, fields='id').execute()
        self.assert_private(file_id, email, ('reader', 'commenter') if role == 'commenter' else ('reader',))

    def share(self, item):
        email = item['email']
        role = item.get('file_role', 'reader')
        if role not in ('reader', 'commenter'):
            raise InputError('Permiso de documento no válido.')
        self.assert_private(item['folder_id'], email)
        self.assert_private(item['file_id'], email, ('reader', 'commenter') if role == 'commenter' else ('reader',))
        if role == 'commenter':
            # An inherited reader may be upgraded on the individual document.
            existing = next((p for p in self.permissions(item['file_id'])
                if p.get('type') == 'user' and p.get('emailAddress', '').lower() == email), None)
            if existing and existing.get('role') == 'reader':
                self.drive.permissions().update(fileId=item['file_id'], permissionId=existing['id'], body={'role':'commenter'}, fields='id').execute()
            elif not existing:
                self.drive.permissions().create(fileId=item['file_id'],
                    body={'type':'user', 'emailAddress':email, 'role':'commenter'},
                    sendNotificationEmail=True, fields='id').execute()
        if not any(p.get('type') == 'user' and p.get('emailAddress', '').lower() == email
                   for p in self.permissions(item['folder_id'])):
            self.drive.permissions().create(fileId=item['folder_id'],
                body={'type':'user', 'emailAddress':email, 'role':'reader'},
                sendNotificationEmail=True, fields='id').execute()
